from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

out="/mnt/data/Glass_Clustering_Assignment_Complete_K2_K3_K4.docx"

doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(0.65)
sec.bottom_margin=Inches(0.65)
sec.left_margin=Inches(0.75)
sec.right_margin=Inches(0.75)

# Set default font
for style_name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
    style = doc.styles[style_name]
    style.font.name = "Times New Roman"
    style.font.color.rgb = RGBColor(0,0,0)
doc.styles["Normal"].font.size = Pt(11)

def add_heading(text, size=13):
    p=doc.add_paragraph()
    r=p.add_run(text)
    r.bold=True
    r.font.name="Times New Roman"
    r.font.size=Pt(size)
    r.font.color.rgb=RGBColor(0,0,0)
    p.paragraph_format.space_before=Pt(7)
    p.paragraph_format.space_after=Pt(4)
    return p

def add_bold_label(label, value=""):
    p=doc.add_paragraph()
    r=p.add_run(label)
    r.bold=True
    r.font.color.rgb=RGBColor(0,0,0)
    if value:
        p.add_run(value)
    return p

def add_bullet(label, value):
    p=doc.add_paragraph(style="List Bullet")
    r=p.add_run(label)
    r.bold=True
    r.font.color.rgb=RGBColor(0,0,0)
    p.add_run(value)
    return p

def add_image_placeholder(text):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(4)
    p.paragraph_format.space_after=Pt(8)
    r=p.add_run(f"\n[ PASTE {text} HERE ]\n")
    r.bold=True
    r.italic=True
    r.font.size=Pt(10)
    r.font.color.rgb=RGBColor(0,0,0)
    return p

# Title
p=doc.add_paragraph()
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("CS-415 Data Warehousing & Data Mining")
r.bold=True; r.font.size=Pt(15); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)
p=doc.add_paragraph()
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("Lab Assignment 4 – Clustering Using WEKA")
r.bold=True; r.font.size=Pt(14); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)

doc.add_paragraph("Name: ______________________________")
doc.add_paragraph("Roll No.: ___________________________")
doc.add_paragraph("Branch/Group: _______________________")
doc.add_paragraph("Date: _______________________________")

add_heading("1. Objective")
doc.add_paragraph(
    "To perform unsupervised learning using clustering techniques in WEKA. "
    "SimpleKMeans is applied for K = 2, K = 3 and K = 4, the resulting clusters are compared "
    "and visualized, and another clustering method is applied for comparison."
)

add_heading("2. Dataset Overview")
add_bold_label("Dataset Name: ", "Glass")
add_bold_label("Number of Instances: ", "214")
add_bold_label("Number of Attributes: ", "10")
doc.add_paragraph(
    "The Glass dataset contains 214 instances and 10 attributes. The attributes include "
    "RI, Na, Mg, Al, Si, K, Ca, Ba, Fe and Type. The numerical measurements are used to form "
    "clusters using the clustering algorithms in WEKA."
)
add_image_placeholder("YOUR GLASS DATASET / PREPROCESS SCREENSHOT")

add_heading("3. Experimental Setup & Results")

# K2
add_heading("1. SimpleKMeans with K = 2", 12)
add_bold_label("Iterations: ", "9")
add_bold_label("Within Cluster Sum of Squared Errors (SSE): ", "118.20374073549189")
add_bold_label("Cluster Centroids:")
add_bullet("Cluster 0: ", "RI = 1.5186, Na = 13.2822, Mg = 3.5483, Al = 1.1718, Si = 72.575, K = 0.4412, Ca = 8.7949, Ba = 0.0118, Fe = 0.0564 (Dominant Type: build wind float)")
add_bullet("Cluster 1: ", "RI = 1.5182, Na = 13.4956, Mg = 2.0813, Al = 1.6356, Si = 72.704, K = 0.5360, Ca = 9.0702, Ba = 0.2890, Fe = 0.0575 (Dominant Type: build wind non-float)")
add_bold_label("Instance Distribution:")
add_bullet("Cluster 0: ", "88 instances (41%)")
add_bullet("Cluster 1: ", "126 instances (59%)")
add_bold_label("SimpleKMeans K=2 Output")
add_image_placeholder("K=2 WEKA OUTPUT SCREENSHOT")
add_bold_label("Cluster Visualization for K=2")
add_image_placeholder("K=2 CLUSTER VISUALIZATION SCREENSHOT")

# K3
add_heading("2. SimpleKMeans with K = 3", 12)
add_bold_label("Iterations: ", "7")
add_bold_label("Within Cluster Sum of Squared Errors (SSE): ", "77.12426898755668")
add_bold_label("Cluster Centroids:")
add_bullet("Cluster 0: ", "RI = 1.5187, Na = 13.1245, Mg = 2.9314, Al = 1.4310, Si = 72.5814, K = 0.4990, Ca = 9.1571, Ba = 0.0467, Fe = 0.0828 (Dominant Type: build wind non-float)")
add_bullet("Cluster 1: ", "RI = 1.5186, Na = 13.2686, Mg = 3.5514, Al = 1.1618, Si = 72.6005, K = 0.4385, Ca = 8.7951, Ba = 0.0124, Fe = 0.0515 (Dominant Type: build wind float)")
add_bullet("Cluster 2: ", "RI = 1.5172, Na = 14.2533, Mg = 0.4916, Al = 2.0260, Si = 72.8902, K = 0.6077, Ca = 8.8681, Ba = 0.7526, Fe = 0.0156 (Dominant Type: headlamps)")
add_bold_label("Instance Distribution:")
add_bullet("Cluster 0: ", "87 instances (41%)")
add_bullet("Cluster 1: ", "84 instances (39%)")
add_bullet("Cluster 2: ", "43 instances (20%)")
add_bold_label("SimpleKMeans K=3 Output")
add_image_placeholder("K=3 WEKA OUTPUT SCREENSHOT")
add_bold_label("Cluster Visualization for K=3")
add_image_placeholder("K=3 CLUSTER VISUALIZATION SCREENSHOT")

# K4
add_heading("3. SimpleKMeans with K = 4", 12)
add_bold_label("Iterations: ", "8")
add_bold_label("Within Cluster Sum of Squared Errors (SSE): ", "75.33518239462457")
add_bold_label("Cluster Centroids:")
add_bullet("Cluster 0: ", "RI = 1.5200, Na = 13.4247, Mg = 3.6169, Al = 0.9706, Si = 72.3166, K = 0.3216, Ca = 9.1688, Ba = 0.0075, Fe = 0.1334 (Dominant Type: build wind non-float)")
add_bullet("Cluster 1: ", "RI = 1.5178, Na = 13.1908, Mg = 3.5152, Al = 1.2704, Si = 72.7494, K = 0.5040, Ca = 8.5787, Ba = 0.0154, Fe = 0.0063 (Dominant Type: build wind float)")
add_bullet("Cluster 2: ", "RI = 1.5172, Na = 14.2533, Mg = 0.4916, Al = 2.0260, Si = 72.8902, K = 0.6077, Ca = 8.8681, Ba = 0.7526, Fe = 0.0156 (Dominant Type: headlamps)")
add_bullet("Cluster 3: ", "RI = 1.5187, Na = 13.1136, Mg = 2.9290, Al = 1.4364, Si = 72.5968, K = 0.5028, Ca = 9.1491, Ba = 0.0467, Fe = 0.0797 (Dominant Type: build wind non-float)")
add_bold_label("Instance Distribution:")
add_bullet("Cluster 0: ", "32 instances (15%)")
add_bullet("Cluster 1: ", "52 instances (24%)")
add_bullet("Cluster 2: ", "43 instances (20%)")
add_bullet("Cluster 3: ", "87 instances (41%)")
add_bold_label("SimpleKMeans K=4 Output")
add_image_placeholder("K=4 WEKA OUTPUT SCREENSHOT")
add_bold_label("Cluster Visualization for K=4")
add_image_placeholder("K=4 CLUSTER VISUALIZATION SCREENSHOT")

# EM
add_heading("4. Expectation-Maximization (EM) Clustering", 12)
doc.add_paragraph(
    "The EM clustering algorithm was applied to the same Glass dataset. The values below will be "
    "filled after running EM in WEKA."
)
add_bold_label("Number of clusters selected by WEKA: ", "______________________________")
add_bold_label("Iterations: ", "______________________________")
add_bold_label("Log likelihood: ", "______________________________")
add_bold_label("Instance Distribution:")
for c in range(4):
    add_bullet(f"Cluster {c}: ", "______________________________")
add_bold_label("EM Clustering Output")
add_image_placeholder("EM WEKA OUTPUT SCREENSHOT")

# Comparison
add_heading("5. Clustering Comparison")
doc.add_paragraph(
    "The results obtained using SimpleKMeans for different values of K are compared using the "
    "number of iterations, within-cluster SSE and the distribution of instances."
)
table=doc.add_table(rows=1, cols=5)
table.style="Table Grid"
table.alignment=WD_TABLE_ALIGNMENT.CENTER
headers=["Algorithm","No. of Clusters","Cluster Size / Distribution","Performance Metric","Iterations"]
for i,h in enumerate(headers):
    cell=table.rows[0].cells[i]
    cell.text=h
    for p in cell.paragraphs:
        for r in p.runs:
            r.bold=True
            r.font.color.rgb=RGBColor(0,0,0)

rows=[
    ["SimpleKMeans","2","C0: 88 (41%), C1: 126 (59%)","SSE: 118.20374073549189","9"],
    ["SimpleKMeans","3","C0: 87 (41%), C1: 84 (39%), C2: 43 (20%)","SSE: 77.12426898755668","7"],
    ["SimpleKMeans","4","C0: 32 (15%), C1: 52 (24%), C2: 43 (20%), C3: 87 (41%)","SSE: 75.33518239462457","8"],
    ["EM","To be recorded","To be recorded","Log likelihood: To be recorded","To be recorded"],
]
for row in rows:
    cells=table.add_row().cells
    for i,v in enumerate(row):
        cells[i].text=v
        cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

# Interpretation
add_heading("6. Interpretation & Discussion")
doc.add_paragraph(
    "For K = 2, SimpleKMeans divided the 214 instances into two clusters. Cluster 0 contained "
    "88 instances (41%) and Cluster 1 contained 126 instances (59%). The within-cluster SSE was "
    "118.20374073549189."
)
doc.add_paragraph(
    "For K = 3, the dataset was divided into three clusters containing 87, 84 and 43 instances. "
    "The SSE decreased to 77.12426898755668. Cluster 2 was mainly associated with the headlamps "
    "type, while the other two clusters were mainly associated with build wind float and build wind non-float."
)
doc.add_paragraph(
    "For K = 4, the four clusters contained 32, 52, 43 and 87 instances. The SSE decreased further "
    "to 75.33518239462457. The headlamps group remained as a separate cluster, while the build wind "
    "categories were divided into additional groups."
)
doc.add_paragraph(
    "The SSE decreases when the number of clusters is increased from 2 to 3 and from 3 to 4. "
    "The cluster sizes also change as the data is divided into more groups. The final comparison with "
    "EM will be added after its WEKA output is obtained."
)

# Conclusion
add_heading("7. Conclusion")
doc.add_paragraph(
    "SimpleKMeans clustering was successfully performed on the Glass dataset using K = 2, K = 3 and K = 4. "
    "The cluster distributions, centroids and within-cluster SSE values were recorded from WEKA, and the "
    "clusters were visualized. The SSE values obtained were 118.20374073549189 for K = 2, "
    "77.12426898755668 for K = 3 and 75.33518239462457 for K = 4. EM clustering was also selected for "
    "comparison and its final results will be added after running the algorithm."
)

doc.save(out)
print(out)
